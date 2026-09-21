"""Write DataFrames to .xlsx files that are ready to filter in Excel."""

from __future__ import annotations

from pathlib import Path
from typing import Dict, Optional

import pandas as pd
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter


def write_table_xlsx(
    df: pd.DataFrame,
    path: Path,
    sheet_name: str = "data",
    number_formats: Optional[Dict[str, str]] = None,
    column_fills: Optional[Dict[str, str]] = None,
    freeze: str = "B2",
) -> Path:
    """One sheet with filter dropdowns on the header, frozen header row / first column,
    column widths fitted to content, optional per-column number formats
    (e.g. {"or_high": "0.00"}) and background colors as hex RGB (e.g. {"or_high": "F2F2F2"},
    applied to the header and every data cell of that column)."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name=sheet_name, index=False)
        ws = writer.sheets[sheet_name]
        ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = freeze
        formats = number_formats or {}
        fills = {col: PatternFill(fill_type="solid", start_color=rgb, end_color=rgb)
                 for col, rgb in (column_fills or {}).items()}
        for i, col in enumerate(df.columns, start=1):
            letter = get_column_letter(i)
            sample = df[col].head(200).map(lambda v: len(f"{v:.2f}") if isinstance(v, float) else len(str(v)))
            width = max([len(str(col))] + sample.tolist()) + 2
            ws.column_dimensions[letter].width = min(width, 40)
            fmt = formats.get(col)
            if fmt:
                for cell in ws[letter][1:]:
                    cell.number_format = fmt
            fill = fills.get(col)
            if fill:
                for cell in ws[letter]:
                    cell.fill = fill
    return path
